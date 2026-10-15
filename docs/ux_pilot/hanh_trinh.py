# -*- coding: utf-8 -*-
"""Hai hành trình của hai người dùng mới (TASK 11). Xem chay_pilot.py."""

import glob
import os

import openpyxl

NHAN_SU = {
    "ten": "Nhân sự — chia nhân viên vào dự án",
    "nganh": "company",
    "nhan": {"person": ("Nhân viên", "Employee"), "slot": ("Dự án", "Project"),
             "session": ("Quý", "Quarter"), "group": ("Cấp bậc", "Level"), "score": ("KPI", "KPI")},
    "max_choices": "3",
    "priority": "score+tier", "tier_mode": "strict", "score_scope": "per_person",
    "tiers": "Senior\nMid\nJunior",
    "bo": "bo_cong_ty",
}
SU_KIEN = {
    "ten": "Hội nghị Chuyển đổi số 2026",
    "nganh": "event",
    "nhan": {"person": ("Khách tham dự", "Attendee"), "slot": ("Workshop", "Workshop"),
             "session": ("Phiên", "Track"), "group": ("Hạng vé", "Ticket"), "score": ("Điểm", "Score")},
    "max_choices": "2",
    "priority": "tier", "tier_mode": "strict", "score_scope": "per_slot",
    "tiers": "VIP\nMember\nGuest",
    "bo": "bo_su_kien",
}


def sang_the(page, the):
    page.click('button[data-tab="%s"]' % the)
    page.wait_for_timeout(900)


def vi_tri_ngoai(thu_muc, cho_nao):
    return glob.glob(os.path.join(thu_muc, cho_nao))


def tao_workspace(page, chup, kb):
    """Bảng chọn -> Mới -> 4 bước. Ghi lại từng màn."""
    chup("00_mo_phan_mem_lan_dau")
    page.click("#btnWorkspace")
    page.wait_for_selector("#wsBtnMoi")
    chup("01_bang_chon_workspace")
    page.click("#wsBtnMoi")
    page.wait_for_selector("#wzTen")
    chup("02_wizard_buoc1_mac_dinh")
    page.fill("#wzTen", kb["ten"])
    page.select_option("#wzNganh", kb["nganh"])
    chup("03_wizard_buoc1_da_dien")
    page.click("#wzTiep")
    page.wait_for_selector("#wzNhan_person_vi")
    chup("04_wizard_buoc2_nhan_mac_dinh")
    for dt, (vi, en) in kb["nhan"].items():
        page.fill("#wzNhan_%s_vi" % dt, vi)
        page.fill("#wzNhan_%s_en" % dt, en)
    chup("05_wizard_buoc2_da_dien")
    page.click("#wzTiep")
    page.wait_for_selector("#wzMaxChoices")
    chup("06_wizard_buoc3_luat_mac_dinh")
    page.fill("#wzMaxChoices", kb["max_choices"])
    # Chọn cách dùng nhóm và phạm vi điểm TRƯỚC rồi mới chọn cách xếp ưu tiên (thứ tự người dùng thật hay làm).
    page.select_option("#wzTierMode", kb["tier_mode"])
    page.select_option("#wzScoreScope", kb["score_scope"])
    page.select_option("#wzPriority", kb["priority"])
    page.fill("#wzTiers", kb["tiers"])
    chup("07_wizard_buoc3_da_dien")
    page.click("#wzTiep")
    page.wait_for_selector(".ws-xem-lai")
    chup("08_wizard_buoc4_xem_lai")
    page.click("#wzTao")
    page.wait_for_function(
        "n => document.querySelector('#wsTenHienTai') && document.querySelector('#wsTenHienTai').textContent === n",
        arg=kb["ten"], timeout=15000)
    page.wait_for_selector("#dropZone")
    page.wait_for_timeout(800)
    chup("09_sau_khi_tao_the_van_hanh_trong")


def tai_so_mau(page, chup, pm):
    page.click("#btnTaiSoMau")
    page.wait_for_timeout(1200)
    chu = chup("10_tai_so_nhap_mau")
    tep = sorted(glob.glob(os.path.join(pm.xuat, "*.xlsx")))
    kq = {}
    if tep:
        wb = openpyxl.load_workbook(tep[-1])
        kq["so_mau_sheet"] = wb.sheetnames
        kq["so_mau_tieu_de"] = {s: [c.value for c in wb[s][1]] for s in wb.sheetnames[:2]}
        kq["so_mau_huong_dan"] = [r[0].value for r in wb[wb.sheetnames[-1]].iter_rows() if r[0].value]
    kq["thong_bao_tai_so_mau"] = chu
    return kq


def chay_khi_chua_co_du_lieu(page, chup):
    page.click("#btnRun")
    page.wait_for_timeout(1500)
    chup("11_bam_chay_khi_chua_co_du_lieu")


def nhap_so(page, chup, duong, ten):
    page.set_input_files("#fileAny", duong)
    page.wait_for_selector("#importQueue:not([hidden])")
    chup(ten + "_a_da_chon_tep")
    if page.locator("#btnImportAll").is_visible():
        page.click("#btnImportAll")
        page.wait_for_timeout(2500)
        return chup(ten + "_b_sau_khi_nhap")
    return None            # sổ bị từ chối ngay ở bước xem trước: không có nút Nhập


def cac_the(page, chup, thu_tu_dau):
    """Mỗi thẻ một ảnh."""
    ra = {}
    for i, the in enumerate(["results", "fallback", "admin", "scoring"], start=thu_tu_dau):
        nut = page.locator('button[data-tab="%s"]' % the)
        if not nut.is_visible():
            ra[the] = None
            continue
        nut.click()
        page.wait_for_timeout(1000)
        ra[the] = chup("%d_the_%s" % (i, the))
    return ra


def hanh_trinh(page, chup, pm, goc, kb):
    ghi = {}
    tao_workspace(page, chup, kb)
    ghi.update(tai_so_mau(page, chup, pm))
    chay_khi_chua_co_du_lieu(page, chup)
    duong_bo = os.path.join(goc, "du_lieu_test", kb["bo"], "SO_NHAP.xlsx")
    # Sai lầm hay gặp của người mới: nạp nhầm sổ của ngành khác
    khac = "bo_su_kien" if kb["bo"] == "bo_cong_ty" else "bo_cong_ty"
    nhap_so(page, chup, os.path.join(goc, "du_lieu_test", khac, "SO_NHAP.xlsx"), "12_nap_nham_so_nganh_khac")
    page.click("#btnClearQueue")
    nhap_so(page, chup, os.path.join(goc, "du_lieu_test", "SO_NHAP_TEST.xlsx"), "13_nap_nham_so_truong_hoc")
    page.click("#btnClearQueue")
    ghi["nhap"] = nhap_so(page, chup, duong_bo, "14_nap_so_dung")
    page.click("#btnHealthRecheck")
    page.wait_for_timeout(1000)
    chup("15_canh_bao_du_lieu")
    page.click("#btnRun")
    page.wait_for_timeout(3000)
    ghi["chay"] = chup("16_sau_khi_chay")
    sang_the(page, "results")
    chup("17_ket_qua_dau_trang")
    try:
        page.click("#btnThuGonTatCa")        # "Mở rộng tất cả"
        page.wait_for_timeout(800)
    except Exception:
        pass
    ghi["ket_qua_mo_rong"] = chup("18_ket_qua_mo_rong_tat_ca")
    for id_ in ("xuatDauVao", "xuatThayDoi"):
        if page.locator("#" + id_).is_enabled():
            page.check("#" + id_)
    page.click("#btnXuatTuyChon")
    page.wait_for_timeout(2500)
    ghi["xuat"] = chup("19_sau_khi_xuat_excel")
    ghi["tep_da_xuat"] = sorted(os.path.basename(t) for t in glob.glob(os.path.join(pm.xuat, "*")))
    ket = sorted(glob.glob(os.path.join(pm.xuat, "ket_qua*.xlsx")))
    if ket:
        wb = openpyxl.load_workbook(ket[-1])
        ghi["xlsx_ket_qua"] = {}
        for s in wb.sheetnames:
            ghi["xlsx_ket_qua"][s] = [[str(c.value) if c.value is not None else "" for c in r]
                                      for r in wb[s].iter_rows(min_row=1, max_row=14)]
    ghi.update({"cac_the": cac_the(page, chup, 20)})
    sang_the(page, "pipeline")
    chup("30_the_van_hanh_cuoi")
    return ghi


def nhan_su(page, chup, pm, goc):
    return hanh_trinh(page, chup, pm, goc, NHAN_SU)


def su_kien(page, chup, pm, goc):
    return hanh_trinh(page, chup, pm, goc, SU_KIEN)
