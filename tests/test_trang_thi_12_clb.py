# -*- coding: utf-8 -*-
"""Trang thi 12 CLB mỗi buổi (docs/thi_12_clb/index.html).

Kiểm: dữ liệu 60 CLB (12 mỗi buổi, 2 câu mỗi CLB, mã và tên không trùng, đáp án CLB mới đã xáo); trang
khớp tệp sinh ra; câu trợ giúp có nguồn trong biểu mẫu; logic chấm trùng chamTheoCLB của .gs; các tình
huống riêng của trang (không thi, thi CLB không xếp Top) chấm đúng; và các luồng trình duyệt thật
(Chromium): chỉ đề CLB đã chọn được dựng, tick tối đa 5, nộp tự tải phiếu, giáo viên gộp phiếu, bỏ Top 11,
sổ CSV đúng điểm.
"""

import csv
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
from collections import Counter

import pytest

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAU = os.path.join(GOC, "mau_forms_thi_clb")
TRANG = os.path.join(GOC, "docs", "thi_12_clb", "index.html")
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

CO = "Có"
KHONG_BUOI = "Không, em bận buổi này"
KHONG_THI = "Em không thi CLB nào buổi này"
DUNG = "Em không thi thêm CLB nào buổi này"


def _nap(ten_tep, ten_module):
    spec = importlib.util.spec_from_file_location(ten_module, os.path.join(GOC, "tests", ten_tep))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def _co_playwright():
    if shutil.which("node") is None or not os.path.exists(CHROME):
        return False
    r = subprocess.run(["node", "-e", "require('playwright')"], capture_output=True, cwd=GOC)
    return r.returncode == 0


def _du_lieu():
    sys.path.insert(0, MAU)
    import de_thi_12_clb as dt
    return dt.du_lieu()


def _nhan(c):
    return "%s (%s)" % (c["name"], c["id"])


def _sai(q):
    return next(o for o in q["opts"] if o != q["dung"])


def _hang(text, ma_hs):
    for row in csv.reader(io.StringIO(text)):
        if ma_hs in row:
            return row
    raise AssertionError("không có dòng " + ma_hs)


def test_12_clb_moi_buoi_tong_60():
    du = _du_lieu()
    assert len(du) == 5
    assert all(len(d["clubs"]) == 12 for d in du), [len(d["clubs"]) for d in du]
    assert sum(len(d["clubs"]) for d in du) == 60
    ids = [c["id"] for d in du for c in d["clubs"]]
    ten = [c["name"] for d in du for c in d["clubs"]]
    assert len(set(ids)) == 60, "trùng mã CLB"
    assert len(set(ten)) == 60, "trùng tên CLB"


def test_moi_clb_hai_cau_bon_lua_chon_dap_an_trong_lua_chon():
    for d in _du_lieu():
        for c in d["clubs"]:
            assert len(c["de"]) == 2, c["id"]
            for q in c["de"]:
                assert len(q["opts"]) == 4 and len(set(q["opts"])) == 4, q["q"]
                assert q["dung"] in q["opts"], q["q"]


def test_dung_10_clb_nhap_danh_dau_ro_rang():
    du = _du_lieu()
    nhap = [c for d in du for c in d["clubs"] if c["nhap"]]
    goc = [c for d in du for c in d["clubs"] if not c["nhap"]]
    assert len(nhap) == 10 and len(goc) == 50
    assert all(sum(1 for c in d["clubs"] if c["nhap"]) == 2 for d in du)


def test_dap_an_clb_moi_da_xao_khong_tap_trung_o_vi_tri_dau():
    """Đáp án CLB mới được xáo như bộ mẫu. Nếu không, đáp án đúng nằm ở vị trí đầu (16/20 câu) và học sinh
    chọn lựa chọn đầu sẽ được khoảng 80% số điểm của các CLB mới mà không cần biết bài."""
    moi = [q for d in _du_lieu() for c in d["clubs"] if c["nhap"] for q in c["de"]]
    assert len(moi) == 20
    dem = Counter(q["opts"].index(q["dung"]) for q in moi)
    assert dem[0] <= 5, dem  # kỳ vọng 5 câu mỗi vị trí
    assert max(dem.values()) <= 8, dem


def test_tep_trang_khop_ban_sinh():
    sys.path.insert(0, MAU)
    import tao_trang_thi_12_clb as tt
    with open(TRANG, encoding="utf-8") as f:
        assert f.read() == tt.noi_dung()


def test_moi_cau_tro_giup_hien_thi_co_nguon_trong_bieu_mau():
    """Mọi câu của chữ trợ giúp đang hiển thị (TRO_GIUP) phải có trong biểu mẫu Google, trừ câu riêng của
    trang (trần Top 1–10). Kiểm chữ thật sự hiển thị, không chỉ danh sách viết tay."""
    sys.path.insert(0, MAU)
    import tao_google_forms as tg
    import tao_trang_thi_12_clb as tt
    gs = (tg.noi_dung().replace("' + BO_TRONG + '", "(Bỏ trống)")
          .replace("' + KHONG_THI + '", "Em không thi CLB nào buổi này")
          .replace("' + '", ""))
    rieng = {"Trang này chỉ có Top 1 đến Top 10."}
    for doan in tt.TRO_GIUP.values():
        for cau in re.split(r"(?<=\.)\s+", doan):
            if cau not in rieng:
                assert cau in gs, cau


def test_cau_tro_giup_co_nguon_trong_bieu_mau():
    sys.path.insert(0, MAU)
    import tao_google_forms as tg
    import tao_trang_thi_12_clb as tt
    gs = tg.noi_dung()
    for cau in tt.CAU_TRO_GIUP_CO_NGUON:
        assert cau in gs, cau


@pytest.mark.skipif(shutil.which("node") is None, reason="cần Node.js để chạy logic chấm")
def test_logic_cham_giong_chamtheoclb(tmp_path):
    """Logic chấm của trang 12 CLB phải cho cùng hàng học sinh và cảnh báo với .gs,
    trên các kịch bản chỉ dùng 50 CLB của bộ mẫu."""
    sys.path.insert(0, MAU)
    import tao_google_forms as tg
    tsk = _nap("test_skill_tao_google_form.py", "tsk_trang_12")
    gs = os.path.join(MAU, "TAO_GOOGLE_FORM.gs")
    tien_to = tsk.CHAM.split("const d=BUOI[0]")[0]
    harness_gs = tmp_path / "gs.js"
    harness_gs.write_text(tien_to + tsk.HARNESS_GS_DUOI_JS, encoding="utf-8")
    harness_web = tmp_path / "web.js"
    harness_web.write_text(tsk.HARNESS_WEB_JS, encoding="utf-8")
    for ten, ds in tsk._kich_ban_trang_web(tg.du_lieu_gs()).items():
        p = tmp_path / (ten + ".json")
        p.write_text(json.dumps(ds, ensure_ascii=False), encoding="utf-8")
        gs_kq = json.loads(subprocess.run(
            ["node", str(harness_gs), gs, str(tmp_path / "x.json"), str(p)],
            capture_output=True, text=True, encoding="utf-8", check=True).stdout)
        web_kq = json.loads(subprocess.run(
            ["node", str(harness_web), TRANG, str(p)],
            capture_output=True, text=True, encoding="utf-8", check=True).stdout)
        bo_trong = lambda rows: [[x for x in r if x != ""] for r in rows]  # noqa: E731
        assert bo_trong(web_kq["hs"]) == bo_trong(gs_kq["hs"]), ten
        assert web_kq["canh"] == gs_kq["canh"], ten


@pytest.mark.skipif(shutil.which("node") is None, reason="cần Node.js để chạy logic chấm")
def test_logic_12_clb_thi_0_va_thi_khong_xep(tmp_path):
    """Trên dữ liệu 12 CLB: không thi CLB nào thì không có điểm và không có cảnh báo; làm bài một CLB
    không xếp trong Top thì không có điểm và có cảnh báo."""
    tsk = _nap("test_skill_tao_google_form.py", "tsk_12_rieng")
    du = _du_lieu()
    d0 = du[0]
    A, C = d0["clubs"][0], d0["clubs"][2]

    def phieu(ma_hs, ten):
        p = {"student_id": ma_hs, "name": ten, "di-thu_2": CO, "thu_2-top-1": _nhan(A)}
        for d in du[1:]:
            p["di-" + d["ma"]] = KHONG_BUOI
        return p

    khong_thi = phieu("HS 401", "Không thi")
    khong_thi["chon-thu_2"] = []
    khong_thi["thi-thu_2"] = KHONG_THI

    khong_xep = phieu("HS 402", "Thi CLB không xếp")
    khong_xep["chon-thu_2"] = [_nhan(C)]
    khong_xep["thi-thu_2"] = _nhan(C)
    for j, q in enumerate(C["de"]):
        khong_xep["%s-%d" % (C["id"], j + 1)] = q["dung"]
    khong_xep["thi-sau-" + C["id"]] = DUNG

    harness = tmp_path / "web.js"
    harness.write_text(tsk.HARNESS_WEB_JS, encoding="utf-8")
    p = tmp_path / "phieu.json"
    p.write_text(json.dumps([khong_thi, khong_xep], ensure_ascii=False), encoding="utf-8")
    kq = json.loads(subprocess.run(
        ["node", str(harness), TRANG, str(p)],
        capture_output=True, text=True, encoding="utf-8", check=True).stdout)

    hs = {r[0]: r for r in kq["hs"][1:]}
    canh = kq["canh"][1:]
    assert all(x == "" for x in hs["HS 401"][3::2]), "không thi mà có điểm"
    assert not [r for r in canh if r[1] == "HS 401"], "không thi mà có cảnh báo"
    assert all(x == "" for x in hs["HS 402"][3::2]), "không xếp Top mà có điểm"
    msg = "%s: làm bài %s nhưng không chọn CLB này ở Top nào, không tính điểm." % (d0["ten"], C["name"])
    assert any(r[1] == "HS 402" and r[2] == msg for r in canh), canh


XLSX_JS = r"""
const fs = require('fs');
const html = fs.readFileSync(process.argv[2], 'utf8');
const doan = html.split('/*XLSX-BEGIN*/')[1].split('/*XLSX-END*/')[0];
const dungXlsx = new Function(doan + '\nreturn dungXlsx;')();
const bytes = dungXlsx([
  { ten: '1. CLB', dong: [['Tên CLB', 'Buổi', 'Chỉ tiêu'], ['CLB Đàn tranh & <Nhạc>', 'thu_2', 20]] },
  { ten: '2. Học sinh', dong: [['Mã HS', 'Họ tên', 'NV1', 'Điểm 1'],
                                ['HS 101', '  Thí sinh  ', 'CLB A', 10],
                                ['HS 102', '=x', 'CLB B', '']] },
]);
fs.writeFileSync(process.argv[3], Buffer.from(bytes));
"""


@pytest.mark.skipif(shutil.which("node") is None, reason="cần Node.js để chạy bộ ghi .xlsx")
def test_bo_ghi_xlsx_doc_duoc_bang_openpyxl(tmp_path):
    """Bộ ghi .xlsx không thư viện phải tạo tệp mà openpyxl đọc đúng từng ô, kể cả ký tự XML và chữ bắt đầu bằng '='."""
    openpyxl = pytest.importorskip("openpyxl")
    js = tmp_path / "xlsx.js"
    js.write_text(XLSX_JS, encoding="utf-8")
    duong = tmp_path / "mau.xlsx"
    subprocess.run(["node", str(js), TRANG, str(duong)], check=True, cwd=GOC, timeout=120)
    wb = openpyxl.load_workbook(duong)
    assert wb.sheetnames == ["1. CLB", "2. Học sinh"]
    assert list(wb["1. CLB"].iter_rows(values_only=True)) == [
        ("Tên CLB", "Buổi", "Chỉ tiêu"), ("CLB Đàn tranh & <Nhạc>", "thu_2", 20)]
    rows = list(wb["2. Học sinh"].iter_rows(values_only=True))
    assert rows[1] == ("HS 101", "  Thí sinh  ", "CLB A", 10)
    assert rows[2] == ("HS 102", "=x", "CLB B", None)


def _kiem_so_nhap_xlsx(duong, A, B):
    """Tệp xuất phải đọc được bằng openpyxl, và bộ kiểm sổ nhập của kiosk không báo lỗi."""
    import openpyxl
    sys.path.insert(0, GOC)
    import so_nhap
    wb = openpyxl.load_workbook(duong)
    assert wb.sheetnames == ["1. CLB", "2. Học sinh"]
    rows1 = list(wb["1. CLB"].iter_rows(values_only=True))
    assert rows1[0] == ("Tên CLB", "Buổi", "Chỉ tiêu", "Suất ưu tiên", "Nhóm ưu tiên", "Mã CLB")
    assert len(rows1) == 61 and all(r[2] == 20 for r in rows1[1:])
    rows2 = list(wb["2. Học sinh"].iter_rows(values_only=True))
    hang = next(r for r in rows2 if r[0] == "HS 101")
    assert hang[1] == "=Thí sinh thử"
    assert hang[2:6] == (A["name"], 10, B["name"], 5)
    assert hang[7] is None  # CLB C xếp Top 3 nhưng không làm bài: không có điểm
    ket = so_nhap.doc_so_nhap(wb)
    assert ket["loi"] == [], ket["loi"]


WALK_JS = r"""
const { chromium } = require('playwright');
const fs = require('fs');
const path = require('path');
const url = 'file://' + process.argv[2];
const thuMuc = process.argv[3];
const kb = JSON.parse(process.argv[4]);
(async () => {
  const browser = await chromium.launch({
    executablePath: '/opt/pw-browsers/chromium-1194/chrome-linux/chrome',
    args: ['--no-sandbox'], headless: true,
  });
  const ctx = await browser.newContext({ acceptDownloads: true, viewport: { width: 390, height: 844 } });
  const page = await ctx.newPage();
  const loi = [];
  page.on('pageerror', e => loi.push('pageerror: ' + e.message));
  page.on('console', m => { if (m.type() === 'error') loi.push('console: ' + m.text()); });
  await page.goto(url);
  const out = { loi: loi };
  const tenCacNut = async () => page.evaluate(() => [...document.querySelectorAll('#form input[type=radio]')].map(e => e.name));
  const tick = o => page.locator('input[type=checkbox][data-key="chon-thu_2"][value="' + o + '"]').click();

  await page.fill('input[data-key="student_id"]', 'HS 101');
  // Tên bắt đầu bằng dấu "=" để kiểm sổ CSV không ghi công thức Excel.
  await page.fill('input[data-key="name"]', '=Thí sinh thử');
  await page.click('#tiep');
  await page.locator('input[type=radio][name="di-thu_2"][value="Có"]').check();
  await page.click('#tiep');

  out.topSelects = await page.locator('select[data-key^="thu_2-top-"]').count();
  out.soOpt = (await page.locator('select[data-key="thu_2-top-1"] option').allTextContents()).length;
  out.tickBoxes = await page.locator('input[type=checkbox][data-key="chon-thu_2"]').count();
  await page.locator('select[data-key="thu_2-top-1"]').selectOption({ label: kb.A.label });
  await page.locator('select[data-key="thu_2-top-2"]').selectOption({ label: kb.B.label });
  await page.locator('select[data-key="thu_2-top-3"]').selectOption({ label: kb.ticThem[0] });
  await tick(kb.A.label);
  await tick(kb.B.label);
  // Tick quá 5: CLB thứ sáu bị từ chối, nên chỉ có 5 CLB được tick.
  for (const o of kb.ticThem) await tick(o);
  out.tickSauQua5 = await page.locator('input[type=checkbox][data-key="chon-thu_2"]:checked').count();
  out.thongBaoTick = await page.locator('.err').allTextContents();
  // Bỏ tick ba CLB thừa để còn đúng hai CLB A và B.
  for (const o of kb.ticThem.slice(0, 3)) await tick(o);
  await page.click('#tiep');

  out.luaChonDau = await page.locator('input[type=radio][name="thi-thu_2"]').count();
  await page.locator('input[type=radio][name="thi-thu_2"][value="' + kb.A.label + '"]').check();
  await page.click('#tiep');
  out.nutDeA = [...new Set(await tenCacNut())];
  for (let j = 0; j < 2; j++) {
    await page.locator('input[type=radio][name="' + kb.A.id + '-' + (j + 1) + '"][value="' + kb.A.dap[j] + '"]').check();
  }
  await page.click('#tiep');
  out.luaChonTiepA = await page.locator('#form input[type=radio]').count();
  await page.locator('input[type=radio][name="thi-sau-' + kb.A.id + '"][value="' + kb.B.label + '"]').check();
  await page.click('#tiep');
  out.nutDeB = [...new Set(await tenCacNut())];
  for (let j = 0; j < 2; j++) {
    await page.locator('input[type=radio][name="' + kb.B.id + '-' + (j + 1) + '"][value="' + kb.B.dap[j] + '"]').check();
  }
  await page.click('#tiep');
  out.luaChonTiepB = await page.locator('#form input[type=radio]').count();
  await page.locator('input[type=radio][name="thi-sau-' + kb.B.id + '"][value="Em không thi thêm CLB nào buổi này"]').check();
  await page.click('#tiep');
  for (const ma of ['thu_3', 'thu_4', 'thu_5', 'thu_6']) {
    await page.locator('input[type=radio][name="di-' + ma + '"][value="Không, em bận buổi này"]').check();
    await page.click('#tiep');
  }
  // Nộp: trình duyệt tự tải tệp phiếu.
  const [tai] = await Promise.all([page.waitForEvent('download'), page.click('#nop')]);
  await tai.saveAs(path.join(thuMuc, 'phieu.json'));
  out.thongBao = await page.locator('.ok').first().textContent();
  out.phieuKhoa = Object.keys(JSON.parse(fs.readFileSync(path.join(thuMuc, 'phieu.json'), 'utf8')).phieu || {});

  // Giáo viên: nạp phiếu của học sinh và một tệp có Top 11 (tệp bị sửa tay).
  const top11 = { loai: 'phieu_thi_12_clb', phien: 1, phieu: { student_id: 'HS 303', name: 'Tệp sửa tay', 'di-thu_2': 'Có' } };
  kb.top11.forEach((o, i) => { top11.phieu['thu_2-top-' + (i + 1)] = o; });
  fs.writeFileSync(path.join(thuMuc, 'phieu_top11.json'), JSON.stringify(top11));
  // Học sinh nộp lại sau với đáp án sai, và một bản sao y hệt của tệp đầu. Dù giáo viên chọn bản sau trước,
  // trang phải giữ bản nộp sớm nhất (điểm 10 và 5) và bỏ bản sao.
  const lai = JSON.parse(fs.readFileSync(path.join(thuMuc, 'phieu.json'), 'utf8'));
  lai.nop_luc = new Date(Date.parse(lai.nop_luc) + 3600000).toISOString();
  lai.phieu[kb.A.id + '-1'] = kb.A.sai[0];
  lai.phieu[kb.A.id + '-2'] = kb.A.sai[1];
  fs.writeFileSync(path.join(thuMuc, 'phieu_lai.json'), JSON.stringify(lai));
  fs.copyFileSync(path.join(thuMuc, 'phieu.json'), path.join(thuMuc, 'phieu_ban_sao.json'));
  await page.locator('#gv-tep').setInputFiles([
    path.join(thuMuc, 'phieu_lai.json'), path.join(thuMuc, 'phieu.json'),
    path.join(thuMuc, 'phieu_ban_sao.json'), path.join(thuMuc, 'phieu_top11.json'),
  ]);
  await page.waitForTimeout(300);
  out.gvSo = await page.locator('#gv-so').textContent();
  out.gvBang = await page.locator('#gv-ket-qua').textContent();
  // Xuất .xlsx: bị chặn khi chưa có chỉ tiêu; sau khi nhập đủ thì tải được sổ nhập.
  await page.locator('#gv-tai-xlsx').click();
  out.thieuChiTieu = await page.locator('#gv-so').textContent();
  await page.evaluate(() => document.querySelectorAll('input[data-cap]').forEach(i => {
    i.value = '20';
    i.dispatchEvent(new Event('input', { bubbles: true }));
  }));
  const [sheet] = await Promise.all([page.waitForEvent('download'), page.locator('#gv-tai-xlsx').click()]);
  out.xlsxTen = sheet.suggestedFilename();
  await sheet.saveAs(path.join(thuMuc, 'so_nhap.xlsx'));
  const [csv] = await Promise.all([page.waitForEvent('download'), page.click('#gv-tai-so')]);
  await csv.saveAs(path.join(thuMuc, 'hoc_sinh.csv'));
  out.csv = fs.readFileSync(path.join(thuMuc, 'hoc_sinh.csv'), 'utf8');
  const [clb] = await Promise.all([page.waitForEvent('download'), page.click('#gv-tai-clb')]);
  await clb.saveAs(path.join(thuMuc, 'clb.csv'));
  out.csvClb = fs.readFileSync(path.join(thuMuc, 'clb.csv'), 'utf8');
  out.tran = await page.evaluate(() => document.documentElement.scrollWidth - window.innerWidth);
  console.log(JSON.stringify(out));
  await browser.close();
})().catch(e => { console.error('FAILED ' + e.stack); process.exit(1); });
"""


@pytest.mark.skipif(not _co_playwright(), reason="cần Node.js, Playwright và Chromium có sẵn")
def test_hoc_sinh_chi_thay_de_cua_clb_da_chon(tmp_path):
    d0 = _du_lieu()[0]
    A, B, C = d0["clubs"][0], d0["clubs"][1], d0["clubs"][2]
    kb = {
        "A": {"label": _nhan(A), "id": A["id"], "dap": [q["dung"] for q in A["de"]],
              "sai": [_sai(q) for q in A["de"]]},
        "B": {"label": _nhan(B), "id": B["id"], "dap": [B["de"][0]["dung"], _sai(B["de"][1])]},
        "ticThem": [_nhan(c) for c in d0["clubs"][2:6]],
        "top11": [_nhan(c) for c in d0["clubs"][:11]],
    }
    js = tmp_path / "walk.js"
    js.write_text(WALK_JS, encoding="utf-8")
    r = subprocess.run(["node", str(js), TRANG, str(tmp_path), json.dumps(kb, ensure_ascii=False)],
                       capture_output=True, text=True, encoding="utf-8", cwd=GOC, timeout=300)
    assert r.returncode == 0, r.stderr
    out = json.loads(r.stdout.strip().splitlines()[-1])
    assert out["loi"] == []
    assert out["topSelects"] == 10, "chỉ xếp Top 1-10 mỗi buổi"
    assert out["soOpt"] == 14 and out["tickBoxes"] == 12
    assert out["tickSauQua5"] == 5 and any("tối đa 5" in t for t in out["thongBaoTick"])
    assert out["luaChonDau"] == 3  # hai CLB đã tick và Không thi
    # Chỉ đề của CLB đã chọn được dựng; không có câu nào của CLB C (không làm bài).
    assert out["nutDeA"] == [A["id"] + "-1", A["id"] + "-2"]
    assert out["nutDeB"] == [B["id"] + "-1", B["id"] + "-2"]
    assert out["luaChonTiepA"] == 2  # CLB B và Không thi thêm
    assert out["luaChonTiepB"] == 1  # chỉ còn Không thi thêm
    assert not any(k.startswith(C["id"] + "-") for k in out["phieuKhoa"])
    assert "Đã nộp" in out["thongBao"] and "gửi" in out["thongBao"]
    # Điểm: A đúng cả hai câu (10); B đúng câu 1 và sai câu 2 (5); C xếp Top 3 nhưng không làm bài (không điểm).
    hang = _hang(out["csv"], "HS 101")
    assert hang[2] == A["name"] and hang[3] == "10"
    assert hang[4] == B["name"] and hang[5] == "5"
    assert hang[6] == C["name"] and hang[7] == ""
    # Tệp có Top 11 bị bỏ: sổ không có cột NV11 và trang báo cho giáo viên.
    assert "Top 11" in out["gvSo"]
    assert "Đã nạp 3 phiếu" in out["gvSo"] and "Bỏ 1 tệp trùng nội dung" in out["gvSo"]
    # Tên có dấu "=" vào sổ CSV dưới dạng chữ, có dấu nháy đơn đứng trước.
    assert "'=Thí sinh thử" in out["csv"]
    assert "Thiếu hoặc sai chỉ tiêu" in out["thieuChiTieu"]
    assert out["xlsxTen"] == "SO_NHAP_CLB.xlsx"
    _kiem_so_nhap_xlsx(tmp_path / "so_nhap.xlsx", A, B)
    assert "NV11" not in out["csv"].splitlines()[0]
    assert _hang(out["csv"], "HS 303")
    # Sổ "1. CLB": 60 dòng và tiêu đề.
    rows = list(csv.reader(io.StringIO(out["csvClb"].lstrip("﻿"))))
    assert len(rows) == 61 and rows[0][0] == "Tên CLB"
    assert out["tran"] <= 0


WALK_SAU_JS = r"""
const { chromium } = require('playwright');
const url = 'file://' + process.argv[2];
(async () => {
  const browser = await chromium.launch({
    executablePath: '/opt/pw-browsers/chromium-1194/chrome-linux/chrome',
    args: ['--no-sandbox'], headless: true,
  });
  const page = await browser.newPage();
  const loi = [];
  page.on('pageerror', e => loi.push('pageerror: ' + e.message));
  await page.goto(url);
  const out = { loi: loi };
  const idCua = s => s.match(/\(([^)]+)\)$/)[1];
  // Mã học sinh toàn khoảng trắng không được qua trang đầu.
  await page.fill('input[data-key="student_id"]', '   ');
  await page.fill('input[data-key="name"]', 'Thí sinh chọn sau');
  await page.click('#tiep');
  out.loiMaTrong = await page.locator('.err').allTextContents();
  await page.fill('input[data-key="student_id"]', 'HS 202');
  await page.click('#tiep');
  await page.locator('input[type=radio][name="di-thu_2"][value="Có"]').check();
  await page.click('#tiep');
  const cacOpt = await page.locator('select[data-key="thu_2-top-1"] option').allTextContents();
  const A = cacOpt[2], B = cacOpt[3];
  await page.locator('input[type=checkbox][data-key="chon-thu_2"][value="' + A + '"]').check();
  await page.locator('input[type=checkbox][data-key="chon-thu_2"][value="' + B + '"]').check();
  await page.click('#tiep');
  // Chọn B trước, dù A đã tick nhưng đứng trước B: trang "tiếp theo" sau B không được đưa A ra.
  await page.locator('input[type=radio][name="thi-thu_2"][value="' + B + '"]').check();
  await page.click('#tiep');
  const idB = idCua(B);
  for (const q of [idB + '-1', idB + '-2']) {
    await page.locator('input[type=radio][name="' + q + '"]').first().check();
  }
  await page.click('#tiep');
  out.tiep = await page.locator('#form input[type=radio]').evaluateAll(els => els.map(e => e.value));
  console.log(JSON.stringify(out));
  await browser.close();
})().catch(e => { console.error('FAILED ' + e.stack); process.exit(1); });
"""


@pytest.mark.skipif(not _co_playwright(), reason="cần Node.js, Playwright và Chromium có sẵn")
def test_tiep_theo_chi_hien_clb_dung_sau_va_chan_ma_trong(tmp_path):
    js = tmp_path / "walk_sau.js"
    js.write_text(WALK_SAU_JS, encoding="utf-8")
    r = subprocess.run(["node", str(js), TRANG], capture_output=True, text=True,
                       encoding="utf-8", cwd=GOC, timeout=300)
    assert r.returncode == 0, r.stderr
    out = json.loads(r.stdout.strip().splitlines()[-1])
    assert out["loi"] == []
    assert out["loiMaTrong"] == ["Em cần ghi mã học sinh."]
    assert out["tiep"] == [DUNG]
