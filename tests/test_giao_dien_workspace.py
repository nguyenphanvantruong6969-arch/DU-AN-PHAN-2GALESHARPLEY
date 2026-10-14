"""Màn chọn / ô đổi / trình tạo không gian làm việc, bằng TRÌNH DUYỆT THẬT.

Python không thấy được ba loại lỗi mà tệp này canh:
  1. ô đổi ở thanh bên không đổi tên khi chuyển workspace;
  2. dữ liệu của workspace cũ còn nằm lại trên màn hình sau khi chuyển;
  3. chuỗi i18n đặt sai bảng, hiện ra thành khoá `ws_...` trên nút.

Máy nào không có playwright/Chromium thì bỏ qua cả file.
"""

import os

import pytest

pytest.importorskip("playwright", reason="chưa cài playwright")
from playwright.sync_api import sync_playwright  # noqa: E402

import browser_host  # noqa: E402
import cau_hinh_ws as cfgws  # noqa: E402
import workspace as w  # noqa: E402
from api import PipelineAPI  # noqa: E402

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

from _trinh_duyet import CHROMIUM  # noqa: E402


@pytest.fixture
def goc_ws(tmp_path):
    p = tmp_path / "goc"
    p.mkdir()
    return str(p)


@pytest.fixture
def mo_app(goc_ws, tmp_path):
    """Trả hàm mở trang cho một PipelineAPI ở chế độ workspace."""
    if CHROMIUM is None:
        pytest.skip("khong tim thay Chromium")
    xuat = tmp_path / "xuat"
    xuat.mkdir()
    api = PipelineAPI(None, thu_muc_xuat=str(xuat), thu_muc_goc=goc_ws)
    url = browser_host.serve(api, GOC, "index.html", open_browser=False)
    with sync_playwright() as pw:
        br = pw.chromium.launch(executable_path=CHROMIUM, args=["--no-sandbox"])

        def _mo(ngon_ngu=None):
            page = br.new_page(bypass_csp=True, viewport={"width": 1280, "height": 800})
            loi = []
            page.on("pageerror", lambda e: loi.append(str(e)))
            if ngon_ngu:
                page.add_init_script(
                    f"try {{ localStorage.setItem('rbda_lang', '{ngon_ngu}'); }} catch (e) {{}}")
            page.goto(url)
            return page, loi

        yield api, _mo
        br.close()


def _ten_o_doi(page):
    return page.locator("#wsTenHienTai").inner_text().strip()


def _cho_ten(page, ten):
    page.wait_for_function(
        "(t) => document.getElementById('wsTenHienTai').textContent.trim() === t", arg=ten)


def _nut_hang(page, slug, khoa):
    return page.locator('.ws-hang[data-slug="%s"] button[data-i18n="%s"]' % (slug, khoa))


def _cho_hop_dong(page):
    page.wait_for_selector("#wsOverlay", state="hidden")


def _mo_tu_man_chon(page, slug):
    _nut_hang(page, slug, "ws_btn_open").click()
    _cho_hop_dong(page)


# ------------------------------------------------------------------ #

def test_man_chon_hien_khi_khoi_dong_va_khong_dong_duoc(mo_app, goc_ws):
    w.tao(goc_ws, "Trường A")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    assert _ten_o_doi(page) == "Chưa chọn"
    assert page.locator("#btnWsDong").is_hidden(), "chưa mở workspace nào thì không có lối thoát"
    page.keyboard.press("Escape")
    assert page.locator("#wsOverlay").is_visible()
    assert _nut_hang(page, "truong-a", "ws_btn_open").count() == 1
    chu = page.locator("#wsOverlay").inner_text()
    assert "ws_" not in chu, "hiện ra khoá i18n thay vì câu chữ: %r" % chu
    assert api.db_path is None
    assert not loi, loi


def test_tao_mo_doi_roi_quay_lai_tieu_de_doi_theo(mo_app, goc_ws):
    """Kịch bản chính: tạo một workspace, chuyển sang nó, chuyển lại, và tên ở
    thanh bên đổi theo mỗi lần."""
    w.tao(goc_ws, "Trường A")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")

    _mo_tu_man_chon(page, "truong-a")
    _cho_ten(page, "Trường A")
    assert api.db_path == os.path.join(goc_ws, "workspaces", "truong-a", "app.db")

    # Tạo workspace thứ hai bằng trình tạo 4 bước; tạo xong tự mở.
    page.locator("#btnWorkspace").click()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    page.locator("#btnWsMoi").click()
    page.wait_for_selector("#wsB1:not([hidden])")
    page.fill("#wsTen", "Dự án B")
    page.locator("#btnWsTiep").click()
    page.wait_for_selector("#wsB2:not([hidden])")
    page.locator("#btnWsTiep").click()
    page.wait_for_selector("#wsB3:not([hidden])")
    page.locator("#btnWsTiep").click()
    page.wait_for_selector("#wsB4:not([hidden])")
    page.locator("#btnWsTao").click()
    _cho_hop_dong(page)
    _cho_ten(page, "Dự án B")
    assert api.db_path == os.path.join(goc_ws, "workspaces", "du-an-b", "app.db")

    # Chuyển về workspace đầu.
    page.locator("#btnWorkspace").click()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    assert page.locator("#btnWsDong").is_visible()   # đã có workspace mở → đóng được
    _mo_tu_man_chon(page, "truong-a")
    _cho_ten(page, "Trường A")
    assert api.db_path == os.path.join(goc_ws, "workspaces", "truong-a", "app.db")
    assert [x["slug"] for x in w.danh_sach(goc_ws)] == ["du-an-b", "truong-a"]
    assert not loi, loi


def test_du_lieu_cua_workspace_cu_khong_con_tren_man_hinh(mo_app, goc_ws):
    a = w.tao(goc_ws, "Trường A")
    b = w.tao(goc_ws, "Trường B")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    _mo_tu_man_chon(page, a)
    _cho_ten(page, "Trường A")
    assert api.create_or_update_club("clb_a", "CLB của A", 3, 0, "")["ok"]

    page.locator('.nav-item[data-tab="admin"]').click()
    page.wait_for_function(
        "document.body.innerText.includes('CLB của A')")

    page.locator("#btnWorkspace").click()
    _mo_tu_man_chon(page, b)
    _cho_ten(page, "Trường B")
    # Vẫn ở thẻ Quản lý, và danh sách đã nạp lại từ workspace B (trống).
    page.wait_for_function("!document.body.innerText.includes('CLB của A')")
    assert api.list_clubs_admin()["data"] == []

    page.locator("#btnWorkspace").click()
    _mo_tu_man_chon(page, a)
    page.wait_for_function("document.body.innerText.includes('CLB của A')")
    assert not loi, loi


def test_so_nhap_dang_cho_khong_sang_workspace_khac(mo_app, goc_ws):
    """Hàng chờ nhập giữ sổ dạng base64: nằm lại thì nhập nhầm vào workspace mới."""
    a = w.tao(goc_ws, "Trường A")
    b = w.tao(goc_ws, "Trường B")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    _mo_tu_man_chon(page, a)
    _cho_ten(page, "Trường A")
    page.evaluate("importQueue = [{ ten: 'so.xlsx', b64: 'AAAA', tomTat: {} }]")
    page.locator("#btnWorkspace").click()
    _mo_tu_man_chon(page, b)
    _cho_ten(page, "Trường B")
    assert page.evaluate("importQueue.length") == 0
    assert not loi, loi


def test_xoa_hai_buoc_va_xoa_workspace_dang_mo(mo_app, goc_ws):
    a = w.tao(goc_ws, "Trường A")
    b = w.tao(goc_ws, "Trường B")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    _mo_tu_man_chon(page, a)
    _cho_ten(page, "Trường A")

    page.locator("#btnWorkspace").click()
    nut = _nut_hang(page, b, "ws_btn_delete")
    nut.click()
    page.wait_for_function(
        "(s) => document.querySelector('.ws-hang[data-slug=\"' + s + '\"] .btn-danger')"
        ".classList.contains('is-confirming')", arg=b)
    assert os.path.isdir(os.path.join(goc_ws, "workspaces", b)), "bấm một lần đã xoá"
    nut.click()
    page.wait_for_selector('.ws-hang[data-slug="%s"]' % b, state="detached")
    rac = os.path.join(goc_ws, "workspaces", "_trash")
    assert len(os.listdir(rac)) == 1
    assert any(f.startswith("app.db.bak-")
               for f in os.listdir(os.path.join(rac, os.listdir(rac)[0])))
    assert _ten_o_doi(page) == "Trường A"      # xoá cái khác không đóng cái đang mở

    # Xoá ngay cái đang mở: về lại màn chọn không có lối thoát.
    nut = _nut_hang(page, a, "ws_btn_delete")
    nut.click()
    nut.click()
    page.wait_for_selector('.ws-hang[data-slug="%s"]' % a, state="detached")
    assert _ten_o_doi(page) == "Chưa chọn"
    assert page.locator("#btnWsDong").is_hidden()
    assert page.locator("#wsRong").is_visible()
    assert api.db_path is None
    assert not loi, loi


def test_doi_ten_va_nhan_ban_trong_hop_dat_ten(mo_app, goc_ws):
    a = w.tao(goc_ws, "Trường A")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    _mo_tu_man_chon(page, a)
    _cho_ten(page, "Trường A")

    page.locator("#btnWorkspace").click()
    _nut_hang(page, a, "ws_btn_rename").click()
    page.wait_for_selector("#wsViewTen:not([hidden])")
    page.fill("#wsTenMoi", "Trường A (đã đổi)")
    page.locator("#btnWsTenLuu").click()
    page.wait_for_selector("#wsViewChon:not([hidden])")
    _cho_ten(page, "Trường A (đã đổi)")
    assert [x["slug"] for x in w.danh_sach(goc_ws)] == [a]   # slug giữ nguyên

    _nut_hang(page, a, "ws_btn_duplicate").click()
    page.wait_for_selector("#wsViewTen:not([hidden])")
    assert page.locator("#wsKemDuLieu").is_visible()
    page.fill("#wsTenMoi", "")
    page.locator("#btnWsTenLuu").click()
    assert page.locator("#wsLoiTen").inner_text().strip() != ""   # tên trống bị chặn
    page.fill("#wsTenMoi", "Bản chép")
    page.locator("#btnWsTenLuu").click()
    page.wait_for_selector('.ws-hang[data-slug="ban-chep"]')
    assert not loi, loi


def test_chon_mau_dien_san_moi_buoc_roi_tao(mo_app, goc_ws):
    nguon = w.tao(goc_ws, "Công ty", {
        "industry": "company",
        "labels": {"person": {"vi": "Nhân viên", "en": "Employee"},
                   "slot": {"vi": "Dự án", "en": ""}},
        "rules": {"max_choices": 3, "max_tests": None, "priority": "score+tier",
                  "tier_mode": "strict", "score_scope": "per_person",
                  "tiers": ["Hạng A", "Hạng B"]}})
    api, mo = mo_app
    assert api.save_as_template(nguon, "Mẫu công ty")["ok"]
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    page.locator("#btnWsMoi").click()
    page.wait_for_selector("#wsB1:not([hidden])")
    page.fill("#wsTen", "Công ty khác")
    page.select_option("#wsMau", "mau-cong-ty")
    assert page.input_value("#wsNganh") == "company"
    assert page.input_value("#wsTen") == "Công ty khác"      # tên đã gõ được giữ
    page.locator("#btnWsTiep").click()
    page.wait_for_selector("#wsB2:not([hidden])")
    assert page.input_value("#wsNhan-person-vi") == "Nhân viên"
    assert page.input_value("#wsNhan-person-en") == "Employee"
    assert page.input_value("#wsNhan-slot-vi") == "Dự án"
    assert page.input_value("#wsNhan-slot-en") == ""          # en trống: dùng chữ Việt
    page.locator("#btnWsTiep").click()
    page.wait_for_selector("#wsB3:not([hidden])")
    assert page.input_value("#wsMaxChoices") == "3"
    assert page.is_checked("#wsKhongGioiHan") and page.is_disabled("#wsMaxTests")
    assert page.input_value("#wsUuTien") == "score+tier"
    assert page.input_value("#wsCheDoNhom") == "strict"
    assert page.input_value("#wsPhamViDiem") == "per_person"
    assert page.is_enabled("#wsNhom")
    assert page.input_value("#wsNhom") == "Hạng A\nHạng B"
    # Đổi sang cách "dự trữ" thì ô thứ tự nhóm bị khoá.
    page.select_option("#wsCheDoNhom", "reserve")
    assert page.is_disabled("#wsNhom")
    page.select_option("#wsCheDoNhom", "strict")
    page.fill("#wsNhom", "Hạng A\n\n Hạng C \n")
    page.locator("#btnWsTiep").click()
    page.wait_for_selector("#wsB4:not([hidden])")
    xem_lai = page.locator("#wsXemLai").inner_text()
    assert "Công ty khác" in xem_lai and "Nhân viên / Employee" in xem_lai
    assert "Hạng A › Hạng C" in xem_lai
    page.locator("#btnWsTao").click()
    _cho_hop_dong(page)
    _cho_ten(page, "Công ty khác")

    cfg = cfgws.doc_cau_hinh(api.db_path)
    assert cfg["industry"] == "company"
    assert cfg["labels"]["person"] == {"vi": "Nhân viên", "en": "Employee"}
    assert cfg["labels"]["slot"] == {"vi": "Dự án", "en": "Dự án"}     # en rỗng -> vi
    assert cfg["rules"] == {
        "max_choices": 3, "max_tests": None, "priority": "score+tier",
        "tier_mode": "strict", "score_scope": "per_person",
        "tiers": ["Hạng A", "Hạng C"]}
    assert not loi, loi


def test_trinh_tao_chan_buoc_sai(mo_app, goc_ws):
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    assert page.locator("#wsRong").is_visible()
    page.locator("#btnWsMoi").click()
    page.wait_for_selector("#wsB1:not([hidden])")
    page.locator("#btnWsTiep").click()                  # tên trống
    assert page.locator("#wsB1").is_visible()
    assert page.locator("#wsLoiBuoc").inner_text().strip() != ""
    page.fill("#wsTen", "X")
    page.locator("#btnWsTiep").click()
    page.fill("#wsNhan-slot-vi", "")
    page.locator("#btnWsTiep").click()                  # thiếu tên tiếng Việt
    assert page.locator("#wsB2").is_visible()
    page.fill("#wsNhan-slot-vi", "Chỗ")
    page.locator("#btnWsTiep").click()
    page.fill("#wsMaxChoices", "21")
    page.locator("#btnWsTiep").click()
    assert page.locator("#wsB3").is_visible()
    page.fill("#wsMaxChoices", "4")
    page.select_option("#wsCheDoNhom", "strict")
    page.fill("#wsNhom", "A\nA")
    page.locator("#btnWsTiep").click()
    assert page.locator("#wsB3").is_visible()           # nhóm trùng tên
    page.fill("#wsNhom", "A\nB")
    page.locator("#btnWsLui").click()
    assert page.locator("#wsB2").is_visible()
    assert page.input_value("#wsNhan-slot-vi") == "Chỗ"  # quay lại không mất chữ đã gõ
    page.locator("#btnWsHuy").click()
    page.wait_for_selector("#wsViewChon:not([hidden])")
    assert w.danh_sach(goc_ws) == []
    assert not loi, loi


def test_nhap_zip_tu_man_chon(mo_app, goc_ws, tmp_path):
    a = w.tao(goc_ws, "Nguồn")
    duong = str(tmp_path / "nguon.zip")
    w.xuat_zip(goc_ws, a, duong)
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")
    page.set_input_files("#wsFileZip", duong)
    page.wait_for_selector('.ws-hang[data-slug="nguon-2"]')
    assert not loi, loi


def test_tieng_anh_khong_lot_chu_viet_va_khong_lo_khoa(mo_app, goc_ws):
    w.tao(goc_ws, "Trường A")
    api, mo = mo_app
    page, loi = mo("en")
    page.wait_for_selector("#wsOverlay:not([hidden])")
    chu = page.locator("#wsOverlay").inner_text()
    assert "Choose a workspace" in chu and "Open" in chu and "Delete" in chu
    assert "Chọn không gian" not in chu and "ws_" not in chu
    assert page.locator("#btnWorkspace").inner_text().startswith("Workspace")
    page.locator("#btnWsMoi").click()
    page.wait_for_selector("#wsB1:not([hidden])")
    page.fill("#wsTen", "Name")
    for i in (2, 3):
        page.locator("#btnWsTiep").click()
        page.wait_for_selector("#wsB%d:not([hidden])" % i)
    chu = page.locator("#wsViewTao").inner_text()
    assert "Step 3 of 4" in chu and "ws_" not in chu
    # Đổi ngôn ngữ giữa chừng: vẽ lại bằng tiếng Việt, không mất dữ liệu đã nhập.
    page.fill("#wsMaxChoices", "7")
    page.locator("#wsViewTao .ws-ngon-ngu").click()
    page.wait_for_function("document.getElementById('wsBuocNoi').textContent.includes('Bước')")
    assert page.input_value("#wsMaxChoices") == "7"
    assert not loi, loi


def test_che_do_csdl_don_khong_thay_gi_cua_workspace(api, tmp_path):
    """Chạy bằng một tệp dữ liệu chỉ định: ô đổi ẩn, không có hộp thoại."""
    if CHROMIUM is None:
        pytest.skip("khong tim thay Chromium")
    url = browser_host.serve(api, GOC, "index.html", open_browser=False)
    with sync_playwright() as pw:
        br = pw.chromium.launch(executable_path=CHROMIUM, args=["--no-sandbox"])
        page = br.new_page(bypass_csp=True)
        loi = []
        page.on("pageerror", lambda e: loi.append(str(e)))
        page.goto(url)
        page.wait_for_selector("#dropZone")
        page.wait_for_function("document.body.dataset.appInit === '1'")
        page.wait_for_function(
            "document.getElementById('dbStatusLine').textContent !== ''")
        assert page.locator("#btnWorkspace").is_hidden()
        assert page.locator("#wsOverlay").is_hidden()
        assert not loi, loi
        br.close()


def test_tai_so_nhap_mau_dung_nhan_cua_workspace_dang_mo(mo_app, goc_ws):
    """Nút "Tải sổ nhập mẫu" ở thẻ Vận hành dùng cấu hình của workspace đang mở."""
    from openpyxl import load_workbook

    ct = w.tao(goc_ws, "Công ty", {
        "industry": "company",
        "labels": {"person": {"vi": "Nhân viên", "en": "Employee"},
                   "slot": {"vi": "Dự án", "en": "Project"}},
        "rules": {"max_choices": 4}})
    mac_dinh = w.tao(goc_ws, "Trường")
    api, mo = mo_app
    page, loi = mo()
    page.wait_for_selector("#wsOverlay:not([hidden])")

    def tai_mau():
        page.evaluate("document.getElementById('feedbackImportAll').textContent = ''")
        truoc = set(os.listdir(api.thu_muc_xuat))
        page.locator("#btnTaiSoMau").click()
        page.wait_for_function(
            "document.getElementById('feedbackImportAll').textContent !== ''")
        moi = sorted(set(os.listdir(api.thu_muc_xuat)) - truoc)
        assert len(moi) == 1, moi
        return load_workbook(os.path.join(api.thu_muc_xuat, moi[0])).sheetnames[:2]

    _mo_tu_man_chon(page, ct)
    _cho_ten(page, "Công ty")
    assert tai_mau() == ["1. Dự án", "2. Nhân viên"]

    page.locator("#btnWorkspace").click()
    _mo_tu_man_chon(page, mac_dinh)
    _cho_ten(page, "Trường")
    assert tai_mau() == ["1. CLB", "2. Học sinh"]       # trường học: như trước nay
    assert not loi, loi
