"""Dấu vân tay của một Sổ nhập CLB: mọi thứ người dùng thấy và Excel dùng.

Sổ .xlsx không so byte được: openpyxl ghi giờ tạo vào docProps, nên hai lần
ghi liền nhau đã khác byte. Thay vào đó so MỌI thứ có nghĩa: tên sheet, từng ô
có giá trị, độ rộng cột, ngăn cố định, danh sách thả xuống / kiểm điểm và các
quy tắc tô đỏ. Hai sổ cùng dấu vân tay thì mở ra giống hệt nhau.
"""

import so_nhap

# Ba bộ dữ liệu đại diện: mẫu trống, một buổi có học sinh, nhiều buổi có nhóm
# ưu tiên và điểm. GOLDEN (tests/du_lieu_ws/so_nhap_truoc.json) được sinh từ
# mã TRƯỚC khi có `cfg`; sổ ghi bằng cfg=None / không gian làm việc mặc định
# phải khớp nó.


def _clb(cid, ten, ct=10, buoi="", up=0, nhom=""):
    return {"club_id": cid, "name": ten, "capacity": ct,
            "reserve_capacity": up, "reserve_group": nhom, "buoi": buoi}


def _nv(cid, thi=False, diem=""):
    return {"club_id": cid, "thi": thi, "diem": diem}


def bo_du_lieu():
    mot_buoi = ([_clb("clb_a", "CLB Tin học", 5), _clb("clb_b", "CLB Bóng rổ", 8, up=2,
                                                       nhom="chinh_sach")],
                [{"student_id": "HS01", "name": "An", "reserve_group": "chinh_sach",
                  "nv": [_nv("clb_a", True, "9,5"), _nv("clb_b")]},
                 {"student_id": "HS02", "name": "Bình", "reserve_group": "",
                  "nv": [_nv("clb_b", True, ""), _nv("clb_a", True, 7.0)]}])
    nhieu_buoi = ([_clb("t2a", "Cờ vua", 6, "thu_2"), _clb("t2b", "Vẽ", 6, "thu_2"),
                   _clb("t5a", "Cờ vua", 6, "thu_5", up=1, nhom="khoi10"),
                   _clb("t5b", "Nhạc", 6, "thu_5")],
                  [{"student_id": "HS10", "name": "Chi", "reserve_group": "khoi10",
                    "nv": [_nv("t2a", True, 8), _nv("t5a", True, "6,5"), _nv("t5b")]}])
    return {
        "trong": ((), ()),
        "mot_buoi": mot_buoi,
        "nhieu_buoi": nhieu_buoi,
    }


def dau_van_tay(duong_dan):
    from openpyxl import load_workbook

    wb = load_workbook(duong_dan)
    ra = {"sheets": wb.sheetnames, "chi tiet": {}}
    for ws in wb.worksheets:
        o = {}
        for hang in ws.iter_rows():
            for c in hang:
                if c.value is not None:
                    o["%d,%d" % (c.row, c.column)] = [c.data_type, str(c.value)]
        rong = {k: round(d.width, 2) for k, d in sorted(ws.column_dimensions.items())}
        dv = sorted(
            [d.type, d.operator or "", d.formula1 or "", str(d.sqref),
             str(d.errorStyle or ""), d.error or ""]
            for d in ws.data_validations.dataValidation)
        cf = sorted(
            [str(rng.sqref), [list(r.formula) for r in regles]]
            for rng, regles in ws.conditional_formatting._cf_rules.items())
        ra["chi tiet"][ws.title] = {
            "o": o, "rong": rong, "dv": dv, "cf": cf,
            "freeze": ws.freeze_panes, "dong": ws.max_row, "cot": ws.max_column,
        }
    return ra


def ghi_bo(thu_muc, **kw):
    """Ghi ba sổ bằng `so_nhap.ghi_so_nhap(**kw)`; trả {tên bộ: dấu vân tay}."""
    ra = {}
    for ten, (clubs, hs) in bo_du_lieu().items():
        p = str(thu_muc / (ten + ".xlsx"))
        so_nhap.ghi_so_nhap(p, clubs, hs, **kw)
        ra[ten] = dau_van_tay(p)
    return ra
